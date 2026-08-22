"""
excel_adapter.py — Adaptador universal de fallback: lee planillas Excel/CSV
exportadas manualmente cuando el ERP no tiene integración directa.

Soporta:
- .xlsx (via openpyxl)
- .csv (via csv nativo)

Configuración en tenant.json → erp:
  tipo: "excel"
  archivo_path: str  — ruta al archivo Excel/CSV
  hoja_productos: str — nombre de hoja (default "Productos")
  hoja_stock: str     — nombre de hoja (default "Stock")
  hoja_ventas: str    — nombre de hoja (default "Ventas")
  hoja_pedidos: str   — nombre de hoja (default "Pedidos")
  formato_numero: str — "cl" (1.234,56) o "en" (1,234.56). Default "cl".
  delimitador_csv: str — para CSV. Default ";"
"""
import csv
import os
from typing import Optional, List, Dict, Any

from core.erp_adapter import ERPAdapter, Producto, Stock, Venta, Pedido
from core.logger import get_logger

log = get_logger("excel_adapter")


def _leer_excel(ruta: str, hoja: str) -> List[Dict[str, Any]]:
    """Lee una hoja de Excel .xlsx y devuelve lista de dicts."""
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError(
            "openpyxl no está instalado. Ejecuta: pip install openpyxl"
        ) from e

    wb = openpyxl.load_workbook(ruta, data_only=True)
    if hoja not in wb.sheetnames:
        raise ValueError(f"Hoja '{hoja}' no encontrada en {ruta}. Hojas disponibles: {wb.sheetnames}")
    ws = wb[hoja]

    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {}
        for h, v in zip(headers, row):
            d[h] = str(v).strip() if v is not None else ""
        rows.append(d)
    return rows


def _leer_csv(ruta: str, delimitador: str = ";") -> List[Dict[str, Any]]:
    """Lee un archivo CSV y devuelve lista de dicts."""
    with open(ruta, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=delimitador)
        rows = []
        for row in reader:
            rows.append({k.strip(): (v or "").strip() for k, v in row.items()})
        return rows


def _leer_hoja(ruta: str, hoja: str, delimitador: str = ";") -> List[Dict[str, Any]]:
    ext = os.path.splitext(ruta)[1].lower()
    if ext == ".csv":
        return _leer_csv(ruta, delimitador)
    return _leer_excel(ruta, hoja)


class ExcelAdapter(ERPAdapter):
    """Adaptador fallback para archivos Excel/CSV locales."""

    def __init__(self, config: dict):
        super().__init__(config)
        self._ruta = config["archivo_path"]
        self._hoja_productos = config.get("hoja_productos", "Productos")
        self._hoja_stock = config.get("hoja_stock", "Stock")
        self._hoja_ventas = config.get("hoja_ventas", "Ventas")
        self._hoja_pedidos = config.get("hoja_pedidos", "Pedidos")
        self._fmt = config.get("formato_numero", "cl")
        self._delim = config.get("delimitador_csv", ";")

        if not os.path.isfile(self._ruta):
            raise FileNotFoundError(f"Archivo no encontrado: {self._ruta}")

    @staticmethod
    def _str(row: dict, *keys: str) -> str:
        for k in keys:
            v = row.get(k, "").strip()
            if v:
                return v
        return ""

    def _float(self, row: dict, *keys: str) -> float:
        for k in keys:
            raw = row.get(k, "").strip()
            if not raw:
                continue
            try:
                if self._fmt == "cl":
                    limpio = raw.replace(".", "").replace(",", ".")
                else:
                    limpio = raw.replace(",", "")
                return float(limpio)
            except (ValueError, TypeError):
                continue
        return 0.0

    @staticmethod
    def _bool(row: dict, *keys: str, default: bool = True) -> bool:
        for k in keys:
            v = row.get(k, "").strip().lower()
            if v in ("false", "0", "no", "inactivo", "inactive"):
                return False
            if v in ("true", "1", "si", "sí", "activo", "active", "yes"):
                return True
        return default

    def test_conexion(self) -> bool:
        try:
            rows = _leer_hoja(self._ruta, self._hoja_productos, self._delim)
            ok = len(rows) > 0
            log.info("test_conexion: %s (%d filas en '%s')", "OK" if ok else "SIN DATA", len(rows), self._hoja_productos)
            return ok
        except Exception as e:
            log.error("test_conexion falló: %s", e)
            return False

    def get_productos(self) -> List[Producto]:
        rows = _leer_hoja(self._ruta, self._hoja_productos, self._delim)
        productos = []
        for row in rows:
            codigo = self._str(row, "codigo", "Codigo", "CODIGO", "code", "Code", "SKU")
            if not codigo:
                continue
            productos.append(Producto(
                codigo=codigo,
                descripcion=self._str(row, "descripcion", "Descripcion", "DESCRIPCION", "name", "Name", "Nombre") or codigo,
                marca=self._str(row, "marca", "Marca", "MARCA", "brand", "Brand") or None,
                precio=self._float(row, "precio", "Precio", "PRECIO", "price", "Price") or None,
                activo=self._bool(row, "activo", "Activo", "ACTIVO", "active", "Active"),
            ))
        log.info("get_productos: %d artículos", len(productos))
        return productos

    def get_stock(self) -> List[Stock]:
        rows = _leer_hoja(self._ruta, self._hoja_stock, self._delim)
        stock = []
        for row in rows:
            codigo = self._str(row, "codigo", "Codigo", "CODIGO", "code", "Code", "SKU")
            if not codigo:
                continue
            stock.append(Stock(
                codigo=codigo,
                bodega=self._str(row, "bodega", "Bodega", "BODEGA", "warehouse", "Warehouse", "Sucursal") or "Principal",
                cantidad=self._float(row, "cantidad", "Cantidad", "CANTIDAD", "qty", "quantity", "Quantity", "Stock"),
            ))
        log.info("get_stock: %d líneas", len(stock))
        return stock

    def get_ventas(self, desde: str, hasta: str) -> List[Venta]:
        rows = _leer_hoja(self._ruta, self._hoja_ventas, self._delim)
        ventas = []
        for row in rows:
            fecha = self._str(row, "fecha", "Fecha", "FECHA", "date", "Date")[:10]
            if not fecha:
                continue
            if not (desde <= fecha <= hasta):
                continue
            codigo = self._str(row, "codigo", "Codigo", "CODIGO", "code", "Code", "SKU")
            if not codigo:
                continue
            ventas.append(Venta(
                fecha=fecha,
                codigo=codigo,
                descripcion=self._str(row, "descripcion", "Descripcion", "DESCRIPCION", "name", "Name") or codigo,
                cantidad=self._float(row, "cantidad", "Cantidad", "CANTIDAD", "qty", "quantity"),
                precio=self._float(row, "precio", "Precio", "PRECIO", "price", "Price", "monto", "Monto"),
                vendedor=self._str(row, "vendedor", "Vendedor", "VENDEDOR", "seller", "Seller") or None,
                cliente=self._str(row, "cliente", "Cliente", "CLIENTE", "client", "Client", "rut") or None,
            ))
        log.info("get_ventas (%s→%s): %d líneas", desde, hasta, len(ventas))
        return ventas

    def get_pedidos(self) -> List[Pedido]:
        rows = _leer_hoja(self._ruta, self._hoja_pedidos, self._delim)
        pedidos_map: Dict[str, Pedido] = {}
        for row in rows:
            numero = self._str(row, "numero", "Numero", "NUMERO", "number", "Number", "id", "Id", "ID")
            if not numero:
                continue
            if numero not in pedidos_map:
                pedidos_map[numero] = Pedido(
                    numero=numero,
                    fecha=self._str(row, "fecha", "Fecha", "FECHA", "date", "Date")[:10],
                    proveedor=self._str(row, "proveedor", "Proveedor", "PROVEEDOR", "supplier", "Supplier") or "Sin proveedor",
                    lineas=[],
                )
            codigo = self._str(row, "codigo", "Codigo", "CODIGO", "code", "Code", "SKU")
            if codigo:
                pedidos_map[numero].lineas.append({
                    "codigo": codigo,
                    "descripcion": self._str(row, "descripcion", "Descripcion", "DESCRIPCION", "name", "Name"),
                    "cantidad": self._float(row, "cantidad", "Cantidad", "CANTIDAD", "qty", "quantity"),
                })
        pedidos = list(pedidos_map.values())
        log.info("get_pedidos: %d órdenes", len(pedidos))
        return pedidos
